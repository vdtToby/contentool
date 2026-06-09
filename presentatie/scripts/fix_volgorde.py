#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Zet de herordende presentatie netjes recht:
 - paginanummers (footers) hernummeren naar de werkelijke slidepositie
 - de agenda-slide (slide 2) laten aansluiten op de nieuwe verteltvolgorde
Verder wordt niets aan de inhoud gewijzigd."""
import re
from pptx import Presentation
from pptx.util import Emu, Inches

SRC = "/root/.claude/uploads/a5676142-f970-5876-a2f5-acd986393432/287f8938-VDT_Klantreis_Presentatie____Hersteld.pptx"
DST = "/home/user/contentool/presentatie/VDT_Klantreis_Presentatie.pptx"

p = Presentation(SRC)

def set_run0(shape, new_text):
    """Vervang de tekst van de eerste run en verwijder eventuele extra runs,
    zodat opmaak (lettertype/kleur/grootte) behouden blijft."""
    para = shape.text_frame.paragraphs[0]
    if not para.runs:
        return
    para.runs[0].text = new_text
    for extra in para.runs[1:]:
        extra._r.getparent().remove(extra._r)

# --- 1. Footers hernummeren naar positie -----------------------------------
fixed = 0
for idx, slide in enumerate(p.slides):
    pos = idx + 1
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        # alleen de paginanummerbox rechtsonder
        if re.fullmatch(r"\d{2}", t) and sh.top > Inches(6.7) and sh.left > Inches(10):
            set_run0(sh, f"{pos:02d}")
            fixed += 1
print("paginanummers bijgewerkt:", fixed)

# --- 2. Agenda-slide (slide 2) op de nieuwe volgorde inrichten --------------
# Nieuwe verteltvolgorde in 4 blokken:
#   1. De kernboodschap            (slide 3)
#   2. Zo werkt het                (slides 4-9: ingang, rollen, Legal MRI, BaseNet, succescriteria)
#   3. De gedachte erachter        (slides 10-11: waarom + filosofie)
#   4. De vijf fases               (slides 12-17)
agenda = {
    "De gedachte":     ("De kernboodschap",   "Waar we voor staan en waar we naartoe bouwen."),
    "De fases":        ("Zo werkt het",       "Eén ingang, de rollen, de Legal MRI en BaseNet."),
    "De rolverdeling": ("De gedachte erachter","Waarom we dit doen en onze vijf principes."),
    "Het systeem":     ("De vijf fases",      "De klantreis stap voor stap."),
}
old_desc = {
    "Waarom we dit doen en waar we naartoe bouwen.": "Waar we voor staan en waar we naartoe bouwen.",
    "Vijf stappen die op elkaar voortbouwen.":       "Eén ingang, de rollen, de Legal MRI en BaseNet.",
    "Wie doet wat — en wie is waarvan eigenaar.":     "Waarom we dit doen en onze vijf principes.",
    "Hoe we dit borgen in BaseNet.":                  "De klantreis stap voor stap.",
}
slide2 = p.slides[1]
changed = 0
for sh in slide2.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text.strip()
    if t in agenda:
        set_run0(sh, agenda[t][0])
        changed += 1
    elif t in old_desc:
        set_run0(sh, old_desc[t])
        changed += 1
print("agenda-elementen bijgewerkt:", changed)

p.save(DST)
print("opgeslagen:", DST, "| slides:", len(p.slides._sldIdLst))
