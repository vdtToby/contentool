#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genereert een professionele PowerPoint-presentatie voor het klantteam van VDT
over de nieuwe klantreis "Van Dossier naar Ondernemerschap".

Bronnen:
  - VDT_Klantreis_V2.docx
  - Conceptmails_onboarding_nieuwe_klanten_alle_fases.docx
Visuele identiteit gebaseerd op vdt-advocaten.nl (merkkleuren uit de huisstijl).

Output: presentatie/VDT_Klantreis_Presentatie.pptx
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------------------
# Huisstijl / merkkleuren VDT
# ---------------------------------------------------------------------------
TEAL    = RGBColor(0x00, 0x7F, 0x81)   # primair
GREEN   = RGBColor(0x2F, 0xA7, 0x66)   # positief / groei
RED     = RGBColor(0xE7, 0x40, 0x49)   # aandacht / accent
YELLOW  = RGBColor(0xF4, 0xC2, 0x00)   # highlight
DARKTEAL= RGBColor(0x00, 0x42, 0x44)   # diep / afgeleid
NAVY    = RGBColor(0x14, 0x2A, 0x33)   # tekst donker

INK     = RGBColor(0x1F, 0x29, 0x37)   # body tekst
SUBINK  = RGBColor(0x52, 0x5C, 0x6B)   # secundaire tekst
LIGHT   = RGBColor(0xF4, 0xF6, 0xF7)   # lichte achtergrond
CARD    = RGBColor(0xFF, 0xFF, 0xFF)   # kaart wit
LINE     = RGBColor(0xE2, 0xE7, 0xEA)  # lijnen
SOFTTEAL = RGBColor(0xE6, 0xF1, 0xF1)  # zachte teal vlakken
SOFTGREEN= RGBColor(0xE9, 0xF5, 0xEE)
SOFTYELL = RGBColor(0xFD, 0xF6, 0xDC)
SOFTRED  = RGBColor(0xFB, 0xE8, 0xE9)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Inter"          # huisstijl-font (valt terug op Calibri indien niet aanwezig)
FONT_H = "Inter"

# Faseschema – kleurgecodeerd door de hele deck
PHASE_COLORS = [TEAL, GREEN, YELLOW, RED, DARKTEAL]
PHASE_SOFT   = [SOFTTEAL, SOFTGREEN, SOFTYELL, SOFTRED, SOFTTEAL]

# ---------------------------------------------------------------------------
# Presentatie-canvas (16:9)
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def slide():
    return prs.slides.add_slide(BLANK)

def _no_line(shp):
    shp.line.fill.background()

def _fill(shp, color):
    shp.fill.solid()
    shp.fill.fore_color.rgb = color

def rect(s, x, y, w, h, color=None, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE,
         shadow=False, radius=None):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    if color is None:
        sp.fill.background()
    else:
        _fill(sp, color)
    # witte kaarten krijgen automatisch een subtiele rand zodat ze ook
    # zonder schaduw-rendering duidelijk afgebakend zijn
    if line is None and shadow and color == CARD:
        line, line_w = LINE, Pt(1)
    if line is None:
        _no_line(sp)
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    if shadow:
        _shadow(sp)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp

def _shadow(sp):
    """Subtiele zachte schaduw onder een kaart."""
    spPr = sp._element.spPr
    el = spPr.makeelement(qn('a:effectLst'), {})
    sh = el.makeelement(qn('a:outerShdw'), {
        'blurRad': '90000', 'dist': '38100', 'dir': '5400000', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '14323B'})
    alpha = clr.makeelement(qn('a:alpha'), {'val': '22000'})
    clr.append(alpha)
    sh.append(clr)
    el.append(sh)
    spPr.append(el)

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         wrap=True, space_after=None, line_spacing=None):
    """runs = list of paragraphs; each paragraph = list of (txt, size, bold, color, font) tuples
       OR a simple (txt, size, bold, color) -> font defaults."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    for m in (tf.margin_left, ):
        pass
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if space_after is not None:
            p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        if line_spacing is not None:
            p.line_spacing = line_spacing
        if isinstance(para, tuple):
            para = [para]
        for r in para:
            run = p.add_run()
            txt, size, bold, color = r[0], r[1], r[2], r[3]
            fnt = r[4] if len(r) > 4 else FONT
            run.text = txt
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = fnt
    return tb

def bg(s, color=WHITE):
    rect(s, 0, 0, SW, SH, color)

def sidebar(s, color, w=Inches(0.18)):
    rect(s, 0, 0, w, SH, color)

def page_header(s, eyebrow, title, accent=TEAL, n=None):
    """Standaard inhoud-koptekst."""
    sidebar(s, accent)
    # accent tab
    rect(s, Inches(0.7), Inches(0.62), Inches(0.55), Inches(0.12), accent)
    text(s, Inches(0.7), Inches(0.78), Inches(11.8), Inches(0.35),
         [[(eyebrow.upper(), 12.5, True, accent)]])
    text(s, Inches(0.7), Inches(1.12), Inches(11.9), Inches(0.95),
         [[(title, 30, True, NAVY)]], line_spacing=1.0)
    if n is not None:
        footer(s, n)

def footer(s, n):
    text(s, Inches(0.7), Inches(7.06), Inches(6), Inches(0.3),
         [[("VDT Klantreis  ·  Van Dossier naar Ondernemerschap", 9, False, SUBINK)]])
    text(s, Inches(11.4), Inches(7.06), Inches(1.4), Inches(0.3),
         [[(f"{n:02d}", 9, True, SUBINK)]], align=PP_ALIGN.RIGHT)

def chip(s, x, y, label, fill, txtcolor=WHITE, w=None, h=Inches(0.34), size=11):
    if w is None:
        w = Inches(0.16 + 0.105 * len(label))
    c = rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x, y - Inches(0.01), w, h, [[(label, size, True, txtcolor)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w

def badge_num(s, cx, cy, num, color, d=Inches(0.62), txtcolor=WHITE, size=20):
    x = cx - d/2
    y = cy - d/2
    rect(s, x, y, d, d, color, shape=MSO_SHAPE.OVAL, shadow=True)
    text(s, x, y, d, d, [[(str(num), size, True, txtcolor)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def icon_badge(s, x, y, glyph, color, d=Inches(0.7), glyphcolor=WHITE, size=22):
    rect(s, x, y, d, d, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3, shadow=True)
    text(s, x, y, d, d, [[(glyph, size, True, glyphcolor)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def bullet_block(s, x, y, w, items, size=14, color=INK, gap=6, marker="—",
                 mcolor=None, bold_first=False, lh=1.08):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            label, sub = it
        else:
            label, sub = it, None
        paras.append([(f"{marker}  ", size, True, mcolor or TEAL),
                      (label, size, bold_first, color)])
    return text(s, x, y, w, Inches(0.4), paras, space_after=gap, line_spacing=lh)

# ===========================================================================
# SLIDE 1 — Titel
# ===========================================================================
s = slide()
bg(s, NAVY)
# achtergrond accentvlakken
rect(s, Inches(0), Inches(0), Inches(0.55), SH, TEAL)
rect(s, Inches(0.55), Inches(0), Inches(0.12), SH, GREEN)
# grote zachte cirkels (decoratief)
rect(s, Inches(9.6), Inches(-1.6), Inches(5.2), Inches(5.2), DARKTEAL, shape=MSO_SHAPE.OVAL)
rect(s, Inches(11.0), Inches(3.4), Inches(4.2), Inches(4.2), TEAL, shape=MSO_SHAPE.OVAL)
rect(s, Inches(10.2), Inches(5.6), Inches(2.2), Inches(2.2), GREEN, shape=MSO_SHAPE.OVAL)

text(s, Inches(1.0), Inches(1.5), Inches(8.6), Inches(0.4),
     [[("VDT ADVOCATEN  ·  KLANTTEAM", 14, True, YELLOW)]])
text(s, Inches(1.0), Inches(2.05), Inches(9.2), Inches(2.4),
     [[("De VDT Klantreis", 52, True, WHITE)],
      [("Van dossier naar ", 34, True, RGBColor(0xCF,0xE6,0xE6)),
       ("ondernemerschap", 34, True, GREEN)]], line_spacing=1.04)
# scheidingslijn
rect(s, Inches(1.05), Inches(4.55), Inches(2.2), Inches(0.06), YELLOW)
text(s, Inches(1.0), Inches(4.8), Inches(8.8), Inches(1.0),
     [[("Wij bouwen geen verzameling losse contactmomenten,", 17, False, RGBColor(0xD7,0xE4,0xE4))],
      [("maar een structureel systeem dat ondernemers beter leert kennen.", 17, False, RGBColor(0xD7,0xE4,0xE4))]],
     line_spacing=1.15)
text(s, Inches(1.0), Inches(6.55), Inches(9), Inches(0.4),
     [[("Voor: klantteam · directie · betrokken advocaten", 12.5, True, RGBColor(0x9F,0xC4,0xC4))]])

# ===========================================================================
# SLIDE 2 — Doel van deze presentatie
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Waarvoor we hier zitten", "Het doel van deze presentatie", TEAL, n=2)

intro = ("In de komende minuten nemen we je mee in de nieuwe klantreis: de gedachte "
         "erachter, de rolverdeling, de workflow in BaseNet en wat dit betekent voor "
         "onze klanten.")
text(s, Inches(0.7), Inches(2.15), Inches(7.0), Inches(1.2),
     [[(intro, 15, False, INK)]], line_spacing=1.25)

cards = [
    ("De gedachte", "Waarom we dit doen en waar we naartoe bouwen.", TEAL),
    ("De fases", "Vijf stappen die op elkaar voortbouwen.", GREEN),
    ("De rolverdeling", "Wie doet wat — en wie is waarvan eigenaar.", YELLOW),
    ("Het systeem", "Hoe we dit borgen in BaseNet.", DARKTEAL),
]
cx = Inches(0.7); cy = Inches(3.7); cw = Inches(2.86); ch = Inches(2.4); gap = Inches(0.18)
for i,(t,d,col) in enumerate(cards):
    x = cx + i*(cw+gap)
    rect(s, x, cy, cw, ch, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
    rect(s, x, cy, cw, Inches(0.12), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    badge_num(s, x+Inches(0.6), cy+Inches(0.78), i+1, col, d=Inches(0.66))
    text(s, x+Inches(0.28), cy+Inches(1.22), cw-Inches(0.56), Inches(0.4),
         [[(t, 15.5, True, NAVY)]])
    text(s, x+Inches(0.28), cy+Inches(1.62), cw-Inches(0.56), Inches(0.7),
         [[(d, 11.5, False, SUBINK)]], line_spacing=1.12)

# toon-strip
tones = ["Praktisch", "Ondernemend", "Klantgericht", "Geen jargon", "Klantwaarde centraal"]
tx = Inches(0.7); ty = Inches(6.4)
for tlabel in tones:
    w = chip(s, tx, ty, tlabel, SOFTTEAL, txtcolor=TEAL, size=11)
    tx += w + Inches(0.16)

# ===========================================================================
# SLIDE 3 — De kernboodschap
# ===========================================================================
s = slide()
bg(s, TEAL)
rect(s, Inches(0), Inches(0), SW, Inches(0.16), YELLOW)
rect(s, Inches(9.8), Inches(-1.4), Inches(5.0), Inches(5.0), DARKTEAL, shape=MSO_SHAPE.OVAL)
rect(s, Inches(11.4), Inches(4.2), Inches(3.6), Inches(3.6), RGBColor(0x00,0x6E,0x70), shape=MSO_SHAPE.OVAL)

text(s, Inches(0.95), Inches(0.85), Inches(8), Inches(0.4),
     [[("DE KERNBOODSCHAP", 13.5, True, YELLOW)]])
text(s, Inches(0.95), Inches(1.55), Inches(10.6), Inches(2.6),
     [[("Het doel is niet meer verkopen.", 30, True, WHITE)],
      [("Het doel is om VDT zó prettig, toegankelijk en waardevol te maken", 24, True, WHITE)],
      [("om mee samen te werken, dat klanten ", 24, True, WHITE),
       ("vanzelf", 24, True, YELLOW),
       (" meer willen afnemen.", 24, True, WHITE)]],
     line_spacing=1.12)

# twee contrastkaarten
y = Inches(4.7); h = Inches(1.85)
rect(s, Inches(0.95), y, Inches(5.5), h, RGBColor(0x00,0x6A,0x6C), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, Inches(1.25), y+Inches(0.22), Inches(5.0), Inches(0.4),
     [[("NIET", 12, True, YELLOW)]])
text(s, Inches(1.25), y+Inches(0.62), Inches(5.0), Inches(1.1),
     [[("Een verzameling losse contactmomenten.", 16, False, WHITE)]], line_spacing=1.15)

rect(s, Inches(6.85), y, Inches(5.5), h, GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, Inches(7.15), y+Inches(0.22), Inches(5.0), Inches(0.4),
     [[("WÉL", 12, True, WHITE)]])
text(s, Inches(7.15), y+Inches(0.62), Inches(5.0), Inches(1.1),
     [[("Eén structureel systeem dat ondernemers beter leert kennen, beter helpt en op het juiste moment meedenkt.", 14.5, True, WHITE)]], line_spacing=1.12)

# ===========================================================================
# SLIDE 4 — Waarom deze klantreis? (dossier vs ondernemer)
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Waarom deze klantreis?", "Van dossierrelatie naar ondernemersrelatie", TEAL, n=4)

# linker kaart: dossierrelatie
lx = Inches(0.7); ty = Inches(2.25); cw = Inches(5.55); ch = Inches(3.45)
rect(s, lx, ty, cw, ch, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, line=LINE, line_w=Pt(1))
icon_badge(s, lx+Inches(0.35), ty+Inches(0.35), "✕", SUBINK, d=Inches(0.6), size=20)
text(s, lx+Inches(1.1), ty+Inches(0.42), cw-Inches(1.3), Inches(0.5),
     [[("De dossierrelatie", 18, True, NAVY)]])
text(s, lx+Inches(1.1), ty+Inches(0.86), cw-Inches(1.3), Inches(0.4),
     [[("Hoe het vaak gaat", 12, True, SUBINK)]])
bullet_block(s, lx+Inches(0.42), ty+Inches(1.5), cw-Inches(0.8),
             ["Contact alleen als er een probleem speelt",
              "Draait om dossiers en losse werkzaamheden",
              "Reactief: wachten tot de klant belt",
              "We kennen het dossier, niet de onderneming"],
             size=13.5, marker="–", mcolor=SUBINK, gap=10)

# rechter kaart: ondernemersrelatie
rx = Inches(7.08)
rect(s, rx, ty, cw, ch, SOFTGREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
rect(s, rx, ty, Inches(0.14), ch, GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
icon_badge(s, rx+Inches(0.35), ty+Inches(0.35), "✓", GREEN, d=Inches(0.6), size=22)
text(s, rx+Inches(1.1), ty+Inches(0.42), cw-Inches(1.3), Inches(0.5),
     [[("De ondernemersrelatie", 18, True, GREEN)]])
text(s, rx+Inches(1.1), ty+Inches(0.86), cw-Inches(1.3), Inches(0.4),
     [[("Waar VDT naartoe bouwt", 12, True, RGBColor(0x2F,0x7A,0x55))]])
bullet_block(s, rx+Inches(0.42), ty+Inches(1.5), cw-Inches(0.8),
             ["Contact vóórdat er een probleem ontstaat",
              "Draait om de ondernemer én de onderneming",
              "Proactief: wij denken mee en kijken vooruit",
              "We begrijpen de onderneming achter het dossier"],
             size=13.5, marker="✓", mcolor=GREEN, gap=10)

# onderschrift / pijl
text(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.6),
     [[("Cross-selling is geen doel op zich — het is het ", 14, False, INK),
       ("logische gevolg", 14, True, TEAL),
       (" van relevantie, vertrouwen en klantwaarde.", 14, False, INK)]],
     align=PP_ALIGN.CENTER)

# ===========================================================================
# SLIDE 5 — De filosofie: 5 principes
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "De gedachte erachter", "De filosofie: vijf principes", GREEN, n=5)

principes = [
    ("We leren de ondernemer kennen", "Niet alleen het dossier — de mens achter de vraag.", TEAL),
    ("We leren de onderneming kennen", "Wat speelt er? Waar liggen uitdagingen en ambities?", GREEN),
    ("We voegen eerst waarde toe", "Eerst helpen, dan adviseren. Nooit andersom.", YELLOW),
    ("We luisteren actief", "Niet wachten op feedback — zelf contact opnemen.", RED),
    ("We kijken vooruit", "Niet alleen vandaag, maar de onderneming van morgen.", DARKTEAL),
]
# bovenste rij 3, onderste rij 2 (gecentreerd)
cw = Inches(3.82); ch = Inches(1.95); gap = Inches(0.22)
def princ_card(x, y, i):
    t,d,col = principes[i]
    rect(s, x, y, cw, ch, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
    rect(s, x, y, Inches(0.13), ch, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    badge_num(s, x+Inches(0.62), y+Inches(0.62), i+1, col, d=Inches(0.58), size=18)
    text(s, x+Inches(1.05), y+Inches(0.34), cw-Inches(1.25), Inches(0.62),
         [[(t, 14.5, True, NAVY)]], line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x+Inches(0.42), y+Inches(1.12), cw-Inches(0.7), Inches(0.7),
         [[(d, 12, False, SUBINK)]], line_spacing=1.14)

topy = Inches(2.3)
startx = Inches(0.7)
for i in range(3):
    princ_card(startx + i*(cw+gap), topy, i)
boty = topy + ch + gap
# twee kaarten gecentreerd
start2 = Inches(0.7) + (cw+gap)/2
for j,i in enumerate([3,4]):
    princ_card(start2 + j*(cw+gap), boty, i)

# ===========================================================================
# SLIDE 6 — De vijf fases: tijdlijn
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Het overzicht", "De vijf fases van de klantreis", TEAL, n=6)

phases = [
    ("Welkom &\nOnboarding", "Direct na\nklantwording", "Warme, persoonlijke start"),
    ("Kennismaking &\nLegal MRI", "± 14 dagen", "Onderneming leren kennen"),
    ("Proactief\nwaarde toevoegen", "± 30 dagen", "Helpen zonder tegenprestatie"),
    ("90 dagen\nCheck-in", "± 90 dagen", "Luisteren en verdiepen"),
    ("Onderneem Mee\nGesprek (OMG)", "± 12 maanden", "Vooruitkijken samen"),
]
# tijdlijn-as
line_y = Inches(3.45)
rect(s, Inches(0.95), line_y, Inches(11.45), Inches(0.06), LINE)
n = len(phases)
left = 1.35
right = 12.0
step = (right-left)/(n-1)
for i,(title,moment,sub) in enumerate(phases):
    cx = Inches(left + i*step)
    col = PHASE_COLORS[i]
    # node
    badge_num(s, cx, line_y+Inches(0.03), i+1, col, d=Inches(0.66), size=20)
    # om-en-om boven/onder
    if i % 2 == 0:
        # boven
        cardy = line_y - Inches(1.55)
        # connector
        rect(s, cx-Inches(0.01), cardy+Inches(1.35), Inches(0.03), Inches(0.2), col)
    else:
        cardy = line_y + Inches(0.55)
        rect(s, cx-Inches(0.01), line_y+Inches(0.36), Inches(0.03), Inches(0.2), col)
    cwid = Inches(2.05)
    cx0 = cx - cwid/2
    rect(s, cx0, cardy, cwid, Inches(1.35), PHASE_SOFT[i], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    rect(s, cx0, cardy, cwid, Inches(0.1), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tcol = NAVY
    text(s, cx0+Inches(0.1), cardy+Inches(0.16), cwid-Inches(0.2), Inches(0.62),
         [[(title.replace("\n"," "), 12.5, True, tcol)]], align=PP_ALIGN.CENTER, line_spacing=0.98)
    text(s, cx0+Inches(0.1), cardy+Inches(0.78), cwid-Inches(0.2), Inches(0.3),
         [[(moment.replace("\n"," "), 10.5, True, col if col!=YELLOW else DARKTEAL)]], align=PP_ALIGN.CENTER)
    text(s, cx0+Inches(0.1), cardy+Inches(1.04), cwid-Inches(0.2), Inches(0.3),
         [[(sub, 9.5, False, SUBINK)]], align=PP_ALIGN.CENTER)

text(s, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.5),
     [[("Elke stap bouwt voort op de vorige: ", 14, False, INK),
       ("Welkom → Kennismaken → Waarde toevoegen → Luisteren → Vooruitkijken", 14, True, TEAL)]],
     align=PP_ALIGN.CENTER)

# ===========================================================================
# SLIDES 7-11 — Per fase: detail + kernwoorden uit de mails
# ===========================================================================
phase_details = [
    {
        "n":1, "title":"Welkom & Onboarding", "moment":"Direct na klantwording",
        "doel":"Een warme, persoonlijke en laagdrempelige start van de samenwerking.",
        "wat":[ "Persoonlijke welkomstmail vanuit het klantteam",
                "Klant aanmaken & gegevens verrijken (kort formulier)",
                "Laten zien: een héél team staat klaar",
                "Direct duidelijk: de lijnen staan altijd open"],
        "words":["Samen","Sparren","Meedenken","Lijnen altijd open"],
        "result":"De ondernemer voelt zich welkom, gezien en serieus genomen.",
    },
    {
        "n":2, "title":"Kennismaking & Legal MRI", "moment":"± 14 dagen na klantwording",
        "doel":"De onderneming achter het dossier leren kennen — niet om iets te verkopen.",
        "wat":[ "Informele kennismaking: een kop koffie",
                "Niet over het dossier, wél over de onderneming",
                "Legal MRI vooraf invullen (± 10 min)",
                "Rapport met inzichten als gespreksbasis"],
        "words":["Onderneming achter het dossier","Kop koffie","Inzicht","Kennismaken"],
        "result":"“VDT begrijpt mijn onderneming en heeft oprechte interesse.”",
    },
    {
        "n":3, "title":"Proactief waarde toevoegen", "moment":"± 30 dagen",
        "doel":"Waarde leveren zónder directe tegenprestatie. Eerst helpen, dan adviseren.",
        "wat":[ "Een praktisch hulpmiddel of praktijkgids delen",
                "Bruikbaar en toepasbaar — geen vaktaal",
                "Geen nieuwsbrief, geen verplichting",
                "“Hier kan ik direct iets mee.”"],
        "words":["Waardevol","Praktisch hulpmiddel","Geen verplichting","Geen verkooppraatje"],
        "result":"VDT helpt ook als er géén juridisch vraagstuk speelt.",
    },
    {
        "n":4, "title":"90 dagen Check-in", "moment":"± 90 dagen",
        "doel":"Luisteren en de relatie verdiepen. Niet evalueren, niet verkopen.",
        "wat":[ "Kort, laagdrempelig telefoontje (5–10 min)",
                "Hoe ervaar je de samenwerking?",
                "Wat kunnen we beter doen?",
                "Ontwikkelingen & inzichten vastleggen"],
        "words":["Hoe ervaar je de samenwerking?","Wat kunnen we beter doen?","Geen verkoopgesprek"],
        "result":"De ondernemer ervaart betrokkenheid en oprechte interesse.",
    },
    {
        "n":5, "title":"Onderneem Mee Gesprek (OMG)", "moment":"± 12 maanden",
        "doel":"Het hoogtepunt: vooruitkijken naar de onderneming, niet het dossier.",
        "wat":[ "Goed gesprek over ondernemerschap",
                "Onderneming · mensen · structuur · business",
                "Legal MRI als voorbereiding en verdieping",
                "Sparren en meedenken — geen verkooppresentatie"],
        "words":["Onderneem Mee Gesprek","Toekomst","Kansen","Uitdagingen","Meedenken"],
        "result":"“VDT denkt mee als ondernemer en is een sparringpartner.”",
    },
]

for pd in phase_details:
    s = slide()
    i = pd["n"]-1
    col = PHASE_COLORS[i]
    soft = PHASE_SOFT[i]
    tcol = DARKTEAL if col==YELLOW else col
    bg(s, WHITE)
    sidebar(s, col)
    # kop met grote fase-badge
    badge_num(s, Inches(1.15), Inches(1.0), pd["n"], col, d=Inches(0.9), size=30)
    text(s, Inches(1.75), Inches(0.62), Inches(9.5), Inches(0.35),
         [[(f"FASE {pd['n']}  ·  {pd['moment'].upper()}", 12, True, tcol)]])
    text(s, Inches(1.75), Inches(0.96), Inches(10.4), Inches(0.7),
         [[(pd["title"], 28, True, NAVY)]])
    # doel-balk
    rect(s, Inches(1.75), Inches(1.72), Inches(10.5), Inches(0.62), soft,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    text(s, Inches(2.0), Inches(1.72), Inches(10.1), Inches(0.62),
         [[("Doel:  ", 13, True, tcol), (pd["doel"], 13, False, INK)]],
         anchor=MSO_ANCHOR.MIDDLE)

    # links: wat doen we (kaart)
    ly = Inches(2.7); lh = Inches(3.2)
    rect(s, Inches(0.7), ly, Inches(6.7), lh, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04, shadow=True)
    text(s, Inches(1.05), ly+Inches(0.28), Inches(6.0), Inches(0.4),
         [[("WAT DOEN WE", 12, True, tcol)]])
    bullet_block(s, Inches(1.05), ly+Inches(0.78), Inches(6.0), pd["wat"],
                 size=13.5, marker="●", mcolor=col, gap=11, lh=1.12)

    # rechts: kernwoorden uit de mails
    rx = Inches(7.7); rw = Inches(4.95)
    rect(s, rx, ly, rw, lh, soft, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    text(s, rx+Inches(0.35), ly+Inches(0.28), rw-Inches(0.6), Inches(0.4),
         [[("KERNWOORDEN UIT DE MAILS", 12, True, tcol)]])
    text(s, rx+Inches(0.35), ly+Inches(0.62), rw-Inches(0.6), Inches(0.35),
         [[("De taal die onze waarden laat zien", 10.5, False, SUBINK)]])
    wy = ly + Inches(1.12)
    for w in pd["words"]:
        rect(s, rx+Inches(0.35), wy, rw-Inches(0.7), Inches(0.46), CARD,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25, line=col, line_w=Pt(1.25))
        text(s, rx+Inches(0.7), wy, rw-Inches(1.0), Inches(0.46),
             [[(w, 12.5, True, tcol)]], anchor=MSO_ANCHOR.MIDDLE)
        # kleine marker
        rect(s, rx+Inches(0.52), wy+Inches(0.15), Inches(0.16), Inches(0.16), col, shape=MSO_SHAPE.OVAL)
        wy += Inches(0.52)

    # onderbalk: gewenst resultaat
    rect(s, Inches(0.7), Inches(6.1), Inches(11.95), Inches(0.62), col,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    rcol = DARKTEAL if col==YELLOW else WHITE
    text(s, Inches(1.05), Inches(6.1), Inches(11.3), Inches(0.62),
         [[("Gewenst resultaat:  ", 13, True, rcol), (pd["result"], 13, False if col!=YELLOW else False, rcol)]],
         anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 6+pd["n"])

# ===========================================================================
# SLIDE 12 — Iedere klant komt hetzelfde binnen (segmentatie)
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Eén ingang voor iedereen", "Iedere klant komt hetzelfde binnen", TEAL, n=12)

text(s, Inches(0.7), Inches(2.12), Inches(11.8), Inches(0.6),
     [[("Iedereen krijgt dezelfde onboarding. Binnen twee weken bepalen we de strategische waarde.", 15, False, INK)]])

# proces: onboarding -> 2 criteria -> beslissing
py = Inches(3.0)
# stap 1
def proc_box(x, w, top, title, sub, col, soft):
    rect(s, x, py, w, Inches(1.5), soft, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    rect(s, x, py, w, Inches(0.1), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x+Inches(0.2), py+Inches(0.26), w-Inches(0.4), Inches(0.4),
         [[(top, 10.5, True, (DARKTEAL if col==YELLOW else col))]])
    text(s, x+Inches(0.2), py+Inches(0.58), w-Inches(0.4), Inches(0.5),
         [[(title, 15, True, NAVY)]], line_spacing=1.0)
    text(s, x+Inches(0.2), py+Inches(1.04), w-Inches(0.4), Inches(0.4),
         [[(sub, 11, False, SUBINK)]], line_spacing=1.05)

proc_box(Inches(0.7), Inches(3.0), "STAP 1", "Gelijke onboarding", "Welkomstmail + gegevens vastleggen", TEAL, SOFTTEAL)
# pijl
rect(s, Inches(3.78), py+Inches(0.55), Inches(0.5), Inches(0.4), TEAL, shape=MSO_SHAPE.CHEVRON, radius=0.5)
proc_box(Inches(4.4), Inches(3.0), "STAP 2", "Toets twee criteria", "Binnen 2 weken beoordeeld", GREEN, SOFTGREEN)
rect(s, Inches(7.48), py+Inches(0.55), Inches(0.5), Inches(0.4), GREEN, shape=MSO_SHAPE.CHEVRON, radius=0.5)
proc_box(Inches(8.1), Inches(3.0), "STAP 3", "Bepaal segment", "Strategisch of regulier", DARKTEAL, SOFTTEAL)

# criteria-kaarten
cy = Inches(4.95)
rect(s, Inches(0.7), cy, Inches(5.6), Inches(1.65), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
text(s, Inches(1.0), cy+Inches(0.2), Inches(5.0), Inches(0.4),
     [[("DE TWEE CRITERIA", 11.5, True, GREEN)]])
# criterium 1
icon_badge(s, Inches(1.0), cy+Inches(0.62), "10+", GREEN, d=Inches(0.5), size=13)
text(s, Inches(1.65), cy+Inches(0.62), Inches(4.5), Inches(0.5),
     [[("Meer dan 10 medewerkers", 13.5, True, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)
# criterium 2
rect(s, Inches(1.0), cy+Inches(1.18), Inches(0.5), Inches(0.5), TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
text(s, Inches(1.0), cy+Inches(1.18), Inches(0.5), Inches(0.5),
     [[("DMU", 10, True, WHITE)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(1.65), cy+Inches(1.15), Inches(4.5), Inches(0.55),
     [[("DMU bekend ", 13.5, True, NAVY), ("(DGA / directie / beslisser)", 11.5, False, SUBINK)]],
     anchor=MSO_ANCHOR.MIDDLE)

# uitkomst-kaart
rect(s, Inches(6.55), cy, Inches(6.1), Inches(1.65), GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
text(s, Inches(6.9), cy+Inches(0.22), Inches(5.4), Inches(0.4),
     [[("BEIDE AANWEZIG?", 11.5, True, RGBColor(0xD6,0xF0,0xE1))]])
text(s, Inches(6.9), cy+Inches(0.55), Inches(5.4), Inches(0.6),
     [[("→  Strategische klant", 22, True, WHITE)]])
text(s, Inches(6.9), cy+Inches(1.12), Inches(5.4), Inches(0.4),
     [[("Verwachting: gemiddeld circa 5 strategische klanten per maand.", 11.5, False, WHITE)]])

# ===========================================================================
# SLIDE 13 — Rolverdeling
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Wie doet wat", "De rolverdeling", DARKTEAL, n=13)

roles = [
    ("De Secretaresse", "DATA & KWALITEIT", TEAL, SOFTTEAL,
     ["Data-input", "Klant aanmaken", "Datakwaliteit bewaken"]),
    ("Het Klantteam", "REGIE & RELATIE", GREEN, SOFTGREEN,
     ["Regie over de klantreis", "Uitvoering: Legal MRI, check-ins, OMG", "Relatiebeheer"]),
    ("De Advocaat", "INHOUD & DOSSIER", DARKTEAL, SOFTTEAL,
     ["Inhoud", "Dossier", "Juridisch advies"]),
]
cw = Inches(3.9); gap = Inches(0.18); cy = Inches(2.25); ch = Inches(3.05)
for i,(name,tag,col,soft,items) in enumerate(roles):
    x = Inches(0.7) + i*(cw+gap)
    rect(s, x, cy, cw, ch, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
    rect(s, x, cy, cw, Inches(0.85), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, x, cy+Inches(0.55), cw, Inches(0.3), col)  # vierkante onderkant header
    tcol = DARKTEAL if col==YELLOW else WHITE
    text(s, x+Inches(0.3), cy+Inches(0.16), cw-Inches(0.6), Inches(0.4),
         [[(name, 17, True, tcol)]])
    text(s, x+Inches(0.3), cy+Inches(0.55), cw-Inches(0.6), Inches(0.3),
         [[(tag, 10.5, True, (RGBColor(0xCF,0xE6,0xE6) if col!=YELLOW else DARKTEAL))]])
    bullet_block(s, x+Inches(0.32), cy+Inches(1.18), cw-Inches(0.6), items,
                 size=13, marker="●", mcolor=col, gap=10, lh=1.12)

# kernboodschap balk
by = Inches(5.65)
rect(s, Inches(0.7), by, Inches(11.95), Inches(1.05), NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, Inches(1.1), by+Inches(0.12), Inches(11.0), Inches(0.45),
     [[("De advocaat is eigenaar van de ", 17, True, WHITE),
       ("inhoud", 17, True, YELLOW),
       (".", 17, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(1.1), by+Inches(0.52), Inches(11.0), Inches(0.45),
     [[("Het klantteam is eigenaar van de ", 17, True, WHITE),
       ("relatie", 17, True, GREEN),
       (".", 17, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)

# ===========================================================================
# SLIDE 14 — De Legal MRI
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Het instrument", "De Legal MRI", GREEN, n=14)

# linkerkolom: wat is het
lx=Inches(0.7); ly=Inches(2.25)
rect(s, lx, ly, Inches(5.7), Inches(3.05), SOFTGREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
icon_badge(s, lx+Inches(0.4), ly+Inches(0.4), "✚", GREEN, d=Inches(0.8), size=26)
text(s, lx+Inches(1.4), ly+Inches(0.45), Inches(4.0), Inches(0.5),
     [[("Een juridische gezondheidsscan", 16, True, NAVY)]], line_spacing=1.0, anchor=MSO_ANCHOR.MIDDLE)
text(s, lx+Inches(0.4), ly+Inches(1.45), Inches(5.0), Inches(1.5),
     [[("Een soort juridische röntgenfoto van de onderneming.", 14, False, INK)],
      [("Invultijd: ± 10 minuten.", 13, True, GREEN)],
      [("Resultaat: analyse, rapport en aanbevelingen — besproken tijdens de kennismaking en het OMG.", 13, False, INK)]],
     line_spacing=1.18, space_after=8)

# rechterkolom: wel/niet
rx=Inches(6.7)
rect(s, rx, ly, Inches(2.9), Inches(3.05), SOFTRED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, rx+Inches(0.3), ly+Inches(0.25), Inches(2.4), Inches(0.4),
     [[("GEEN", 13, True, RED)]])
bullet_block(s, rx+Inches(0.3), ly+Inches(0.75), Inches(2.3),
             ["Verkoopinstrument","Werk creëren","Dossiers genereren"],
             size=13, marker="✕", mcolor=RED, gap=12)

rx2=Inches(9.75)
rect(s, rx2, ly, Inches(2.9), Inches(3.05), SOFTGREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, rx2+Inches(0.3), ly+Inches(0.25), Inches(2.4), Inches(0.4),
     [[("WÉL", 13, True, GREEN)]])
bullet_block(s, rx2+Inches(0.3), ly+Inches(0.75), Inches(2.3),
             ["Inzicht geven","Bewustwording","Gesprek verdiepen","Prioriteiten zichtbaar"],
             size=13, marker="✓", mcolor=GREEN, gap=10)

# onderbalk
rect(s, Inches(0.7), Inches(5.7), Inches(11.95), Inches(0.95), GREEN,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
text(s, Inches(1.1), Inches(5.7), Inches(11.2), Inches(0.95),
     [[("“Handig — ik krijg inzicht in mijn onderneming.”", 18, True, WHITE)],
      [("Dát is het gevoel dat de Legal MRI moet oproepen. Geen verkoop, wel inzicht.", 12.5, False, RGBColor(0xE3,0xF5,0xEA))]],
     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)

# ===========================================================================
# SLIDE 15 — BaseNet Workflow
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Hoe we het borgen", "De workflow in BaseNet", TEAL, n=15)

text(s, Inches(0.7), Inches(2.1), Inches(11.8), Inches(0.5),
     [[("We weten inmiddels hoe we deze klantreis volledig kunnen inrichten in BaseNet.", 15, False, INK)]])

feats = [
    ("Automatische taken", "Het juiste moment komt vanzelf naar boven.", TEAL),
    ("Workflow-stappen", "Elke fase is een vaste stap in het systeem.", GREEN),
    ("Segmentatie", "Strategische klanten herkend en gemarkeerd.", YELLOW),
    ("Herinneringen", "Niets valt tussen wal en schip.", RED),
    ("Eigenaarschap", "Voor elke taak is duidelijk wie aan zet is.", DARKTEAL),
]
cw=Inches(2.28); gap=Inches(0.12); cy=Inches(2.85); ch=Inches(2.05)
for i,(t,d,col) in enumerate(feats):
    x=Inches(0.7)+i*(cw+gap)
    rect(s, x, cy, cw, ch, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07, shadow=True)
    rect(s, x, cy, cw, Inches(0.1), col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    icon_badge(s, x+Inches(0.28), cy+Inches(0.3), str(i+1), col, d=Inches(0.55), size=17)
    text(s, x+Inches(0.24), cy+Inches(0.98), cw-Inches(0.48), Inches(0.5),
         [[(t, 13.5, True, NAVY)]], line_spacing=1.0)
    text(s, x+Inches(0.24), cy+Inches(1.42), cw-Inches(0.48), Inches(0.55),
         [[(d, 10.5, False, SUBINK)]], line_spacing=1.1)

# onderbalk: geen geheugen, maar een systeem
by=Inches(5.4)
rect(s, Inches(0.7), by, Inches(5.85), Inches(1.3), SOFTRED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, Inches(1.05), by+Inches(0.2), Inches(5.2), Inches(0.4),
     [[("GEEN", 12, True, RED)]])
text(s, Inches(1.05), by+Inches(0.55), Inches(5.2), Inches(0.6),
     [[("Een proces dat afhankelijk is van geheugen.", 16, True, NAVY)]], line_spacing=1.0)

rect(s, Inches(6.8), by, Inches(5.85), Inches(1.3), TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, Inches(7.15), by+Inches(0.2), Inches(5.2), Inches(0.4),
     [[("WÉL", 12, True, YELLOW)]])
text(s, Inches(7.15), by+Inches(0.55), Inches(5.2), Inches(0.6),
     [[("Een systeem. Betrouwbaar, herhaalbaar en geborgd.", 16, True, WHITE)]], line_spacing=1.0)

# ===========================================================================
# SLIDE 16 — Succesfactor: het systeem
# ===========================================================================
s = slide()
bg(s, LIGHT)
page_header(s, "Waar het om draait", "De kracht zit in het systeem", GREEN, n=16)

text(s, Inches(0.7), Inches(2.1), Inches(11.8), Inches(0.7),
     [[("De kracht zit niet in losse activiteiten — niet in één mail, één telefoontje of één gesprek.", 15, False, INK)],
      [("De kracht zit erin dat elke stap voortbouwt op de vorige.", 15, True, NAVY)]],
     line_spacing=1.25)

# flow keten
flow = [("Welkom",TEAL),("Kennismaken",GREEN),("Waarde toevoegen",YELLOW),("Luisteren",RED),("Vooruitkijken",DARKTEAL)]
fy=Inches(3.5); fh=Inches(0.95)
fx=Inches(0.85); fw=Inches(2.05)
for i,(lab,col) in enumerate(flow):
    x=fx+i*(fw+Inches(0.28))
    rect(s, x, fy, fw, fh, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.16, shadow=True)
    tc = DARKTEAL if col==YELLOW else WHITE
    text(s, x, fy, fw, fh, [[(lab, 14.5, True, tc)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(flow)-1:
        ax = x+fw+Inches(0.02)
        rect(s, ax, fy+Inches(0.3), Inches(0.26), Inches(0.35), SUBINK, shape=MSO_SHAPE.CHEVRON, radius=0.5)

# uitkomst kaart
oy=Inches(4.95)
rect(s, Inches(0.85), oy, Inches(11.6), Inches(1.55), CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
rect(s, Inches(0.85), oy, Inches(0.16), Inches(1.55), GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
text(s, Inches(1.3), oy+Inches(0.24), Inches(10.9), Inches(0.5),
     [[("Het resultaat: iets wat veel advocatenkantoren niet hebben", 16, True, NAVY)]])
text(s, Inches(1.3), oy+Inches(0.72), Inches(10.9), Inches(0.7),
     [[("Een relatie die draait om de ondernemer en de onderneming — in plaats van uitsluitend om juridische dossiers. Dat is de ambitie van VDT.", 14, False, INK)]],
     line_spacing=1.2)

# ===========================================================================
# SLIDE 17 — Succescriteria van de onboarding (audit checklist)
# ===========================================================================
s = slide()
bg(s, WHITE)
page_header(s, "Na 90 dagen", "Succescriteria van de onboarding", GREEN, n=17)

text(s, Inches(0.7), Inches(2.12), Inches(11.8), Inches(0.5),
     [[("De onboarding is geslaagd als de klant:", 15, False, INK)]])

checks = [
    "Exact weet wie hij of zij moet bellen bij VDT.",
    "Minimaal twee mensen binnen VDT persoonlijk kent.",
    "Zich gekend voelt en merkt dat VDT de onderneming écht begrijpt.",
    "Minimaal één keer proactief is geholpen — zonder eigen verzoek.",
]
cx = Inches(0.7); cy = Inches(2.8); cw = Inches(11.95); ch = Inches(3.05)
rect(s, cx, cy, cw, ch, CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04, shadow=True)
rect(s, cx, cy, cw, Inches(0.12), GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
text(s, cx+Inches(0.45), cy+Inches(0.32), cw-Inches(0.9), Inches(0.45),
     [[("AUDIT CHECKLIST", 13, True, GREEN)]])
rowy = cy + Inches(0.86)
rowh = Inches(0.46)
for i, item in enumerate(checks):
    y = rowy + i*(rowh+Inches(0.09))
    # zachte rij-achtergrond
    rect(s, cx+Inches(0.45), y, cw-Inches(0.9), rowh, SOFTGREEN,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.22)
    icon_badge(s, cx+Inches(0.55), y+Inches(0.04), "✓", GREEN, d=Inches(0.42), size=16)
    text(s, cx+Inches(1.2), y, cw-Inches(1.7), rowh,
         [[(item, 14.5, False, NAVY)]], anchor=MSO_ANCHOR.MIDDLE)

# resultaatbalk
ry = Inches(6.1)
rect(s, Inches(0.7), ry, Inches(11.95), Inches(0.95), GREEN,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
text(s, Inches(1.1), ry, Inches(11.2), Inches(0.95),
     [[("Het resultaat:  ", 15, True, WHITE),
       ("een relatie die verder gaat dan een juridisch dossier.", 15, False, WHITE)],
      [("Adoptie van VDT over de volle breedte.", 15, True, RGBColor(0xEF,0xF8,0xF1))]],
     anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1, align=PP_ALIGN.CENTER)

# ===========================================================================
# SLIDE 18 — Afsluiting / quote
# ===========================================================================
s = slide()
bg(s, NAVY)
rect(s, Inches(0), Inches(0), SW, Inches(0.16), GREEN)
rect(s, Inches(0), Inches(7.34), SW, Inches(0.16), YELLOW)
rect(s, Inches(-1.5), Inches(3.0), Inches(5.0), Inches(5.0), DARKTEAL, shape=MSO_SHAPE.OVAL)
rect(s, Inches(10.5), Inches(-1.5), Inches(5.0), Inches(5.0), TEAL, shape=MSO_SHAPE.OVAL)

text(s, Inches(1.2), Inches(1.15), Inches(6), Inches(0.4),
     [[("WANNEER IS DEZE KLANTREIS GESLAAGD?", 13.5, True, YELLOW)]])
# grote aanhalingstekens
text(s, Inches(1.0), Inches(1.3), Inches(3), Inches(1.6),
     [[("“", 120, True, GREEN)]])
text(s, Inches(2.3), Inches(2.55), Inches(9.6), Inches(2.6),
     [[("VDT kent niet alleen mijn juridische dossiers,", 30, True, WHITE)],
      [("maar begrijpt óók mijn onderneming.", 30, True, GREEN)]],
     line_spacing=1.18)
text(s, Inches(2.35), Inches(4.95), Inches(9), Inches(0.4),
     [[("— de ondernemer die wij voor ogen hebben", 14, False, RGBColor(0xAE,0xCB,0xCB))]])

rect(s, Inches(2.35), Inches(5.7), Inches(8.4), Inches(0.9), RGBColor(0x00,0x42,0x44),
     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
text(s, Inches(2.35), Inches(5.7), Inches(8.4), Inches(0.9),
     [[("Dán is deze klantreis geslaagd.", 19, True, YELLOW)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------------------
out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
out = os.path.join(out_dir, "VDT_Klantreis_Presentatie.pptx")
prs.save(out)
print("Opgeslagen:", out)
print("Aantal slides:", len(prs.slides._sldIdLst))
